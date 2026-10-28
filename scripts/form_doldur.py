#!/usr/bin/env python3
"""STP 596 bilgilendirme formunu doldurur.

Kullanım:
  python3 form_doldur.py form.md --sablon assets/stp596_bilgilendirme_form.docx --cikti 514xxxxxx_Soyad_stp596_bilgilendirme_form.docx --tarih 02.12.2026

form.md biçimi references/form.md içinde anlatılmıştır: '# Başlık', sonra '## Yazar', '## Özet',
'## Anahtar Kelimeler', '## Tez Çalışması Durumu', '## Tezin ... ile İlişkisi', '## Seminer Konusu',
'## Seminer – Tez İlişkisi', isteğe bağlı '## Teşekkür', '## Yapay Zekâ Kullanım Beyanı', '## Referanslar'.
İngilizce form: başlıklar İngilizce ('## Abstract', '## Keywords', '## Thesis Status', '## Relation of the
Thesis to ...', '## Seminar Topic', '## Seminar–Thesis Relation', '## Acknowledgements', '## Declaration of
AI Use', '## References') ve --dil en (ya da '## Abstract' görülünce otomatik). Şablonun sabit metinleri
(etiketler, bölüm başlıkları, altbilgi) seçilen dilde yazılır.
Şablonun stilleri ve iki sütunlu düzeni korunur; şablondaki örnek metinler silinir.
Bölüm metninde kendi paragrafındaki '$$ LaTeX $$' satırı Word denklemine çevrilir ve sağda (1), (2)… diye
numaralanır (IEEE); pandoc yoksa yer tutucu yazılır ve uyarılır (scripts/denklem.py).
Gerekli paket: python-docx (pip install python-docx); isteğe bağlı: pandoc.
"""
import argparse, copy, os, re, sys
from docx import Document
from docx.text.paragraph import Paragraph
from lxml import etree
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import denklem

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
EQ_RE = re.compile(r"^\$\$(.+?)\$\$$", re.S)
EQ = {"n": 0, "orta": 2420, "sag": 4840, "uyari": []}   # denklem sayacı, sekme konumları (twip), uyarılar

# ---------- form.md ayrıştırma ----------

def parse_form_md(path):
    text = open(path, encoding="utf-8").read().replace("\r\n", "\n")
    title = ""
    sections = {}   # anahtar: normalize edilmiş başlık -> (orijinal başlık, metin)
    order = []
    cur = None
    for line in text.split("\n"):
        if line.startswith("# ") and not title:
            title = line[2:].strip()
            continue
        m = re.match(r"^##\s+(.*)$", line)
        if m:
            cur = m.group(1).strip()
            sections[cur] = []
            order.append(cur)
            continue
        if cur is not None:
            sections[cur].append(line)
    out = {k: "\n".join(v).strip() for k, v in sections.items()}
    return title, out, order

def norm(s):
    s = s.lower().replace("â", "a").replace("î", "i").replace("û", "u")
    s = s.replace("ı", "i").replace("ş", "s").replace("ç", "c").replace("ğ", "g").replace("ö", "o").replace("ü", "u")
    return re.sub(r"[^a-z0-9]+", " ", s).strip()

def find_section(sections, *keys, prefix=None):
    for k in sections:
        nk = norm(k)
        if prefix and nk.startswith(norm(prefix)):
            return k, sections[k]
        for key in keys:
            if nk == norm(key):
                return k, sections[k]
    return None, None

def paragraphs_of(block):
    """Boş satırla ayrılmış paragraflar; tek satır kırılmaları birleştirilir."""
    if not block:
        return []
    block = re.sub(r"(?m)^[ \t]*(\$\$.+?\$\$)[ \t]*$", r"\n\1\n", block.strip())   # denklem kendi paragrafı
    paras = [re.sub(r"\s*\n\s*", " ", p).strip() for p in re.split(r"\n\s*\n", block.strip())]
    return [p for p in paras if p]

def donem_from_tarih(tarih, dil="tr"):
    """'02.12.2026' → '2026-2027 Güz Yarıyılı' / '2026–2027 Fall Semester'; Şubat–Temmuz → Bahar / Spring."""
    if not tarih: return ""
    m = re.match(r"^\s*(\d{1,2})[./-](\d{1,2})[./-](\d{4})", tarih) or re.match(r"^\s*(\d{4})-(\d{1,2})-(\d{1,2})", tarih)
    if not m: return ""
    g = [int(x) for x in m.groups()]
    ay, yil = (g[1], g[2]) if g[2] > 31 else (g[1], g[0])
    guz = ay >= 8 or ay == 1
    y0 = yil if ay >= 8 else yil - 1
    if dil == "en":
        return f"{y0}–{y0+1} {'Fall' if guz else 'Spring'} Semester"
    return f"{y0}-{y0+1} {'Güz' if guz else 'Bahar'} Yarıyılı"

# ---------- dil ----------

LABELS = {
    "tr": {
        "ozet": "Özet", "anahtar": "Anahtar Kelimeler",
        "h1": ["Tez Çalışması", "Seminer"],
        "h2": ["Tez Çalışması Durumu", None, "Seminer Konusu", "Seminer – Tez İlişkisi"],   # None: form.md başlığı aynen
        "tesekkur": "Teşekkür", "beyan": "Yapay Zekâ Kullanım Beyanı", "referanslar": "Referanslar",
        "altbilgi": "İ.T.Ü. Havacılık Enstitüsü – {PROGRAM} {DÖNEM} STP 596 Bilimsel Araştırma, Etik ve Seminer Dersi",
    },
    "en": {
        "ozet": "Abstract", "anahtar": "Keywords",
        "h1": ["Thesis Work", "Seminar"],
        "h2": ["Thesis Status", None, "Seminar Topic", "Seminar–Thesis Relation"],
        "tesekkur": "Acknowledgements", "beyan": "Declaration of AI Use", "referanslar": "References",
        "altbilgi": "ITU Aviation Institute – {PROGRAM} {DÖNEM} STP 596 Scientific Research, Ethics and Seminar",
    },
}

def detect_dil(sections):
    k, _ = find_section(sections, "Abstract", "Keywords", "References", "Thesis Status")
    return "en" if k else "tr"

# ---------- docx yardımcıları ----------

def set_runs(p, texts_with_rpr):
    """Paragrafın bütün run'larını siler; (metin, rPr-kopyası) listesini yazar."""
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    # hyperlink vb. kalıntıları da temizle
    for child in list(p._p):
        if child.tag in (W + "r", W + "hyperlink"):
            p._p.remove(child)
    for text, rpr in texts_with_rpr:
        r = p.add_run(text)
        if rpr is not None:
            r._r.insert(0, copy.deepcopy(rpr))
    return p

def first_rpr(p, idx=0):
    runs = p.runs
    if len(runs) > idx and runs[idx]._r.rPr is not None:
        return copy.deepcopy(runs[idx]._r.rPr)
    return None

def set_text(p, text, rpr=None):
    if rpr is None:
        rpr = first_rpr(p)
    return set_runs(p, [(text, rpr)])

def clone_after(p):
    new = copy.deepcopy(p._p)
    p._p.addnext(new)
    return Paragraph(new, p._parent)

def remove(p):
    p._p.getparent().remove(p._p)

def eq_tabs(doc):
    """İki sütunlu bölümün sütun genişliğinden denklem satırının orta ve sağ sekmesi (twip)."""
    for sct in doc.sections:
        cols = sct._sectPr.find(W + "cols")
        if cols is not None and cols.get(W + "num") == "2":
            space = int(cols.get(W + "space", "720"))
            col = (sct.page_width.twips - sct.left_margin.twips - sct.right_margin.twips - space) // 2
            return col // 2, col - 40
    return EQ["orta"], EQ["sag"]

def eq_paragraph(latex):
    """IEEE denklem satırı: ortada denklem, sağda (n). Çevrilemezse (None, yer tutucu metni)."""
    EQ["n"] += 1
    n = EQ["n"]
    omml = denklem.latex_to_omml(latex)
    if omml is None:
        neden = "pandoc bulunamadı" if not denklem.pandoc_var() else "LaTeX çevrilemedi"
        EQ["uyari"].append(f"Denklem ({n}) Word denklemine çevrilemedi ({neden}); yer tutucu yazıldı. Word'de Ekle → Denklem ile yazın.")
        return None, f"[ÖĞRENCİ TAMAMLAYACAK: denklem ({n}), Word'de Ekle → Denklem ile: {latex}]"
    w = W[1:-1]
    p = etree.fromstring(
        f'<w:p xmlns:w="{w}"><w:pPr><w:pStyle w:val="BodyText"/><w:tabs><w:tab w:val="clear" w:pos="288"/>'
        f'<w:tab w:val="center" w:pos="{EQ["orta"]}"/><w:tab w:val="right" w:pos="{EQ["sag"]}"/></w:tabs>'
        f'<w:spacing w:before="60" w:after="60"/><w:ind w:left="0" w:firstLine="0"/><w:jc w:val="left"/></w:pPr>'
        f'<w:r><w:tab/></w:r><w:r><w:tab/><w:t>({n})</w:t></w:r></w:p>')
    p[1].addnext(copy.deepcopy(omml))          # sekme run'ından sonra denklem, sonra sekme + numara
    return p, None

def fill_block(template_paras, texts, label=""):
    """İlk şablon paragrafını kalıp alır; texts kadar paragraf yazar, fazla şablon paragraflarını siler.
    '$$ … $$' paragrafı numaralı Word denklemi olur."""
    if not texts:
        texts = [f"[ÖĞRENCİ TAMAMLAYACAK: {label}]"]
    base = template_paras[0]
    rpr = first_rpr(base)
    tmpl = copy.deepcopy(base._p)
    for extra in template_paras[1:]:
        remove(extra)
    last = None
    for i, t in enumerate(texts):
        if i == 0:
            p = base
        else:
            el = copy.deepcopy(tmpl); last._p.addnext(el); p = Paragraph(el, base._parent)
        m = EQ_RE.match(t.strip())
        eq, ph = eq_paragraph(m.group(1).strip()) if m else (None, None)
        if eq is not None:
            p._p.addnext(eq); remove(p); p = Paragraph(eq, base._parent)
        else:
            set_text(p, ph or t, rpr)
            if last is not None and EQ_RE.match(texts[i - 1].strip()) and t[:1].islower():
                # denklemden sonra küçük harfle süren cümle ("burada …", "where …") girintisiz başlar (IEEE)
                ppr = p._p.get_or_add_pPr()
                ind = ppr.find(W + "ind")
                if ind is None:
                    ind = etree.Element(W + "ind")
                    after = [c for c in ppr if c.tag in (W + "jc", W + "rPr", W + "sectPr", W + "pPrChange")]
                    after[0].addprevious(ind) if after else ppr.append(ind)   # şema sırası: ind, jc, rPr'den önce
                ind.set(W + "firstLine", "0")
        last = p
    return last

def body_paras_after(doc, heading_p):
    """Bir başlığın altındaki ardışık Body Text paragraflarını verir."""
    out = []
    el = heading_p._p.getnext()
    while el is not None and el.tag == W + "p":
        p = Paragraph(el, heading_p._parent)
        if p.style is not None and p.style.name == "Body Text":
            out.append(p)
            el = el.getnext()
        else:
            break
    return out

# ---------- ana iş ----------

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("form_md")
    ap.add_argument("--sablon", required=True)
    ap.add_argument("--cikti", required=True)
    ap.add_argument("--donem", help="Altbilgideki dönem yazısı; verilmezse --tarih'ten hesaplanır (Ağu–Oca: Güz, Şub–Tem: Bahar)")
    ap.add_argument("--program", help="Altbilgideki program adı; verilmezse yazar bloğunun 2. satırından alınır")
    ap.add_argument("--tarih", help="Seminer tarihi GG.AA.YYYY; dönem bundan hesaplanır")
    ap.add_argument("--dil", choices=["tr", "en"], help="Çıktı dili; verilmezse form.md başlıklarından (Abstract/Keywords → en) anlaşılır")
    a = ap.parse_args()

    title, sec, order = parse_form_md(a.form_md)
    dil = a.dil or detect_dil(sec)
    L = LABELS[dil]
    doc = Document(a.sablon)
    EQ["orta"], EQ["sag"] = eq_tabs(doc)
    P = doc.paragraphs
    by_style = lambda name: [p for p in doc.paragraphs if p.style is not None and p.style.name == name]
    warnings = []

    # Başlık
    set_text(by_style("paper title")[0], title or "[ÖĞRENCİ TAMAMLAYACAK: seminer başlığı]")

    # Yazar bloğu: satır satır; şablonun satır biçimleri (2. satır italik) korunur
    _, yazar = find_section(sec, "Yazar", "Author")
    lines = [l.strip() for l in (yazar or "").split("\n") if l.strip()]
    if len(lines) < 5:
        warnings.append(f"Yazar bloğunda {len(lines)} satır var; beklenen 5 (ad, program, kurum, şehir, iletişim).")
    ap_ = by_style("Author")[0]
    # şablon satır rPr'leri: run'ları satırlara böl
    tmpl_line_rpr, cur = [], None
    for r in ap_.runs:
        if r._r.findall(W + "br"):
            tmpl_line_rpr.append(cur); cur = None
        elif cur is None and r.text.strip():
            cur = copy.deepcopy(r._r.rPr) if r._r.rPr is not None else None
    tmpl_line_rpr.append(cur)
    parts = []
    for i, l in enumerate(lines):
        rpr = tmpl_line_rpr[i] if i < len(tmpl_line_rpr) else tmpl_line_rpr[-1]
        parts.append((l, rpr))
    set_runs(ap_, [])
    for i, (l, rpr) in enumerate(parts):
        r = ap_.add_run(l)
        if rpr is not None:
            r._r.insert(0, copy.deepcopy(rpr))
        if i < len(parts) - 1:
            ap_.add_run().add_break()

    # Özet: "Özet" (italik) + "—" + metin
    _, ozet = find_section(sec, "Özet", "Ozet", "Abstract")
    ozet_t = " ".join(paragraphs_of(ozet)) if ozet else "[ÖĞRENCİ TAMAMLAYACAK: özet]"
    pa = by_style("Abstract")[0]
    set_runs(pa, [(L["ozet"], first_rpr(pa, 0)), ("—", first_rpr(pa, 1)), (ozet_t, first_rpr(pa, 2))])
    n_words = len(ozet_t.split())
    if not (120 <= n_words <= 280):
        warnings.append(f"Özet {n_words} sözcük; hedef 150–250.")

    # Anahtar kelimeler
    _, ak = find_section(sec, "Anahtar Kelimeler", "Keywords")
    ak_t = re.sub(r"\s*\n\s*", " ", (ak or "").strip())
    pk = by_style("Keywords")[0]
    set_text(pk, L["anahtar"] + "—" + (ak_t or "[ÖĞRENCİ TAMAMLAYACAK]"))

    # Ana başlıklar (Heading 1): dil İngilizceyse çevrilir
    if dil == "en":
        for hp, txt in zip(by_style("Heading 1"), L["h1"]):
            set_text(hp, txt)

    # Alt bölümler (Heading 2 sırası: 1.1, 1.2, 2.1, 2.2)
    h2 = by_style("Heading 2")
    if len(h2) < 4:
        sys.exit("Şablonda dört Heading 2 bekleniyordu.")
    sec_12 = find_section(sec, prefix="Tezin")
    if not sec_12[0]: sec_12 = find_section(sec, prefix="Relation of the Thesis")
    spec = [
        (h2[0], find_section(sec, "Tez Çalışması Durumu", "Tez Calismasi Durumu", "Thesis Status", "Thesis Work Status"), None, "tez durumu"),
        (h2[1], sec_12, "iliski", "tezin ilişkisi"),
        (h2[2], find_section(sec, "Seminer Konusu", "Seminar Topic"), None, "seminer konusu"),
        (h2[3], find_section(sec, "Seminer - Tez İlişkisi", "Seminer – Tez İlişkisi", "Seminer Tez İlişkisi",
                             "Seminar - Thesis Relation", "Seminar – Thesis Relation", "Seminar Thesis Relation", "Seminar–Thesis Relation"), None, "seminer–tez ilişkisi"),
    ]
    texts_11 = texts_22 = []
    for i, (heading, (key, block), mode, label) in enumerate(spec):
        if mode == "iliski" and key:
            set_text(heading, key.strip())          # başlık programa göre değişebilir
        elif dil == "en" and L["h2"][i]:
            set_text(heading, L["h2"][i])
        texts = paragraphs_of(block)
        if not texts:
            warnings.append(f"'{label}' bölümü boş; yer tutucu yazıldı.")
        bodies = body_paras_after(doc, heading)
        fill_block(bodies, texts, label)
        if label == "tez durumu": texts_11 = texts
        if label == "seminer–tez ilişkisi": texts_22 = texts

    # 1.1 ile 2.2 benzerliği
    import difflib
    if texts_11 and texts_22:
        ratio = difflib.SequenceMatcher(None, " ".join(texts_11), " ".join(texts_22)).ratio()
        if ratio > 0.6:
            warnings.append(f"1.1 ile 2.2 metinleri %{ratio*100:.0f} benzer; 2.2 seminerin tezdeki yerini anlatmalı.")

    # Arka madde: Teşekkür (varsa) ve Yapay zekâ beyanı, Referanslar başlığından (Heading 5) önce,
    # IEEE'deki gibi numarasız başlıkla (Heading 5 kalıbı) + Body Text
    _, beyan = find_section(sec, "Yapay Zekâ Kullanım Beyanı", "Yapay Zeka Kullanim Beyani", "Declaration of AI Use", "AI Use Declaration", prefix="Yapay Zek")
    if beyan is None:
        _, beyan = find_section(sec, prefix="Declaration of AI")
    _, tesekkur = find_section(sec, "Teşekkür", "Tesekkur", "Teşekkürler", "Acknowledgements", "Acknowledgments", "Acknowledgement", "Acknowledgment")
    h5 = by_style("Heading 5")[0]
    body_tmpl = body_paras_after(doc, h2[3])[0]
    beyan_paras = paragraphs_of(beyan) or ["[ÖĞRENCİ TAMAMLAYACAK: yapay zekâ kullanım beyanı; kalıp references/yz_beyani.md]"]
    if not beyan:
        warnings.append("Yapay zekâ beyanı yok; yer tutucu yazıldı.")
    blocks = []
    if paragraphs_of(tesekkur):
        blocks.append((L["tesekkur"], paragraphs_of(tesekkur)))
    blocks.append((L["beyan"], beyan_paras))
    for head_txt, paras in blocks:
        new_h = copy.deepcopy(h5._p); h5._p.addprevious(new_h)
        new_h = Paragraph(new_h, h5._parent); set_text(new_h, head_txt)
        last = new_h
        for t in paras:
            nb = copy.deepcopy(body_tmpl._p); last._p.addnext(nb)
            last = Paragraph(nb, h5._parent); set_text(last, t, first_rpr(body_tmpl))
    set_text(h5, L["referanslar"])

    # Referanslar: numaralı liste stili; baştaki [n] atılır
    _, refs = find_section(sec, "Referanslar", "Kaynaklar", "References")
    ref_lines = []
    for l in (refs or "").split("\n"):
        l = l.strip()
        if not l: continue
        if re.match(r"^\[\d+\]", l) or re.match(r"^\d+[.)]", l) or not ref_lines:
            ref_lines.append(re.sub(r"^(\[\d+\]|\d+[.)])\s*", "", l))
        else:
            ref_lines[-1] += " " + l           # satır devamı
    refps = [p for p in by_style("references") if p.text.strip()]
    if not ref_lines:
        warnings.append("Referans yok.")
    fill_block(refps, ref_lines, "referanslar")
    if len(ref_lines) < 3:
        warnings.append(f"{len(ref_lines)} referans var; en az 3 beklenir.")
    for i, r in enumerate(ref_lines, 1):
        if not re.search(r"(19|20)\d\d", r):
            warnings.append(f"Referans [{i}] yıl içermiyor.")

    # Lorem kalıntısı denetimi
    full = "\n".join(p.text for p in doc.paragraphs)
    if "Lorem ipsum" in full or "lorem ipsum" in full:
        warnings.append("Belgede 'Lorem ipsum' kaldı.")
    for mark in re.findall(r"\[(ÖĞRENCİ TAMAMLAYACAK|DOĞRULANACAK)[^\]]*\]", full):
        warnings.append(f"İşaret kaldı: [{mark}…]")

    # Altbilgi: {PROGRAM} ve {DÖNEM} yer tutucuları (eski sabit metinli şablon da tanınır)
    donem = a.donem or donem_from_tarih(a.tarih, dil)
    program = a.program or (lines[1] if len(lines) > 1 else "")
    if not donem:
        warnings.append("Dönem belirlenemedi (--tarih ya da --donem verin); altbilgide {DÖNEM} boş bırakıldı.")
    for sct in doc.sections:
        for part in (sct.footer, sct.first_page_footer, sct.even_page_footer):
            for fp in part.paragraphs:
                if "STP 596" not in fp.text: continue
                t = fp.text
                if dil == "en": t = L["altbilgi"]
                if "{" in t:
                    t = t.replace("{PROGRAM}", program).replace("{DÖNEM}", donem or "")
                    t = re.sub(r"\s{2,}", " ", t).replace("–  ", "– ")
                else:
                    if donem: t = re.sub(r"20\d\d-20\d\d\s+\S+\s+Yarıyılı", donem, t)
                    if program: t = re.sub(r"–\s.*?Programı", "– " + program, t)
                set_text(fp, t)
    warnings += EQ["uyari"]
    doc.save(a.cikti)
    print(f"Yazıldı: {a.cikti}  (dil: {dil})" + (f"; {EQ['n']} denklem" if EQ["n"] else ""))
    for w in warnings:
        print("UYARI:", w)

if __name__ == "__main__":
    main()
