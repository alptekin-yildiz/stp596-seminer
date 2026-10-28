#!/usr/bin/env python3
"""STP 596 form (.docx) ve sunum (.pptx) dosyalarını dersin kurallarına göre denetler.

Kullanım:
  python3 kontrol.py 514xxxxxx_Soyad_stp596_bilgilendirme_form.docx 514xxxxxx_Soyad_stp596_sunum.pptx
  (tek dosya da verilebilir)

Denetimler, SKILL.md'deki kontrol listesinin makineyle bakılabilen kısmıdır:
  pptx: slayt sayısı (15–20 içerik), zorunlu slaytlar (İçindekiler, Teşekkürler, Kaynaklar), punto (<18 uyarı;
        Kaynaklar, alt yazı, altbilgi, sayfa no hariç), madde sayısı (>5), kalan yer tutucular, görsel eksikliği,
        kaynak numarası tutarlılığı ([n] metinde var, Kaynaklar'da yok), konuşma notu eksikliği, yapay zekâ satırı,
        punto tutarlılığı (başlıklar arası, slaytlar arası ve slayt içi gövde puntosu), elle yazılmış matematik,
        görseller (slaytta aynı anda 3'ten çok görsel/grafik, çizildiği boyuttan çok küçültülmüş ya da çok küçük görsel;
        PowerPoint grafik nesnesinde 4'ten çok seri ve 14 pt'den küçük ya da grafikler arası farklı punto).
  docx: bölümlerin doluluğu, Lorem ipsum kalıntısı, 1.1–2.2 benzerliği, özet uzunluğu, referans sayısı ve yıl,
        metinde anılmayan/eksik referans numaraları, yapay zekâ beyanı, kalan yer tutucular, elle yazılmış matematik.
Şekil ve grafiklerin içindeki yazı resim olduğundan denetlenemez; o kısım references/son_tur.md ile gözle bakılır.
Türkçe ve İngilizce çıktıları tanır (bölüm ve slayt adları iki dilde aranır).
Çıkış kodu: 0 sorun yok, 1 uyarı var.
"""
import difflib, re, sys

def tlow(s):
    return s.replace("İ", "i").replace("I", "ı").lower()

# İki dilli anahtar sözcükler (Türkçe ve İngilizce çıktılar); eşleme hem Türkçe hem düz küçük harfle denenir
K_ICINDEKILER = ("içindekiler", "contents", "outline")
K_KAYNAK = ("kaynak", "reference")
K_TESEKKUR = ("teşekkür", "acknowledg", "thank")
K_YARDIMCI = K_ICINDEKILER + K_KAYNAK + K_TESEKKUR
def has_k(title, keys): return any(k in tlow(title) or k in title.lower() for k in keys)
def starts_k(title, keys): return tlow(title).startswith(keys) or title.lower().startswith(keys)

MARK = re.compile(r"\[(ÖĞRENCİ TAMAMLAYACAK|DOĞRULANACAK|ÖĞRENCİ KARARI)[^\]]*\]")

# Denklem nesnesine dönüştürülmemiş, düz metinle yazılmış matematik izleri
MATH = [
    (re.compile(r"(?<![\w@./])[A-Za-zΑ-Ωα-ω]{1,3}_\{?[A-Za-z0-9α-ω∞]{1,6}\}?(?![\w@./])"), "alt çizgiyle indis (x_1)"),
    (re.compile(r"[\w)\]]\^[\w{(\-+]"), "şapkayla üs (x^2)"),
    (re.compile(r"\\(int|iint|oint|sum|prod|frac|sqrt|partial|nabla|infty|cdot|times|approx|leq?|geq?|neq|alpha|beta|gamma|delta|Delta|theta|lambda|mu|rho|sigma|omega|Omega|pi|phi|varphi|eta|epsilon|varepsilon|tau|mathrm|mathbf|left|right)\b"), "LaTeX komutu (\\int)"),
    (re.compile(r"\$[^$\n]{1,80}\$"), "dolar işaretli LaTeX ($...$)"),
    (re.compile(r"\b(sqrt|integral|int|sum|exp|ln|log)\s*\("), "fonksiyon adıyla yazılmış işlem (sqrt(...))"),
    (re.compile(r"(?<=[\w)])\s?\*\s?(?=[\w(])"), "yıldızla çarpma (a*b)"),
    (re.compile(r"(?<![<>=!])(<=|>=|!=)(?![<>=])"), "klavye karşılaştırması (<=, >=, !=)"),
]

def math_hits(text):
    hits = []
    for rx, label in MATH:
        m = rx.search(text)
        if m and "http" not in text[max(0, m.start() - 40):m.end()]:
            hits.append(f"{label}: \"{text[max(0, m.start() - 15):m.end() + 15].strip()}\"")
    return hits

def check_docx(path):
    from docx import Document
    out = []
    d = Document(path)
    paras = [(p.style.name if p.style is not None else "", p.text.strip()) for p in d.paragraphs]
    full = "\n".join(t for _, t in paras)
    if re.search(r"lorem ipsum", full, re.I): out.append("Lorem ipsum kalıntısı var.")
    for sct in d.sections:
        for part in (sct.footer, sct.first_page_footer, sct.even_page_footer):
            for fp in part.paragraphs:
                if "{PROGRAM}" in fp.text or "{DÖNEM}" in fp.text: out.append("Altbilgide doldurulmamış yer tutucu: " + fp.text[:60])
                if re.search(r"20(1\d|2[0-5])[-–]20\d\d", fp.text): out.append("Altbilgide eski dönem yazısı: " + fp.text[:80])
    for m in MARK.finditer(full): out.append(f"İşaret kaldı: {m.group(0)[:60]}")
    # bölümler
    sections, cur = {}, None
    for st, t in paras:
        if st.startswith("Heading") and t:
            cur = t; sections[cur] = []
        elif cur and st == "Body Text" and t:
            sections[cur].append(t)
    is_beyan = lambda h: starts_k(h, ("yapay zek", "declaration of ai", "ai use"))
    is_tesekkur = lambda h: starts_k(h, ("teşekkür", "acknowledg"))
    is_refs = lambda h: starts_k(h, ("referans", "references", "kaynak"))
    h2 = {p.text.strip() for p in d.paragraphs if p.style is not None and p.style.name == "Heading 2"}
    for h, body in sections.items():
        if h not in h2 and not is_beyan(h) and not is_tesekkur(h): continue
        if not body: out.append(f"'{h}' bölümü boş.")
    def find(*prefixes):
        for h, b in sections.items():
            if starts_k(h, tuple(p.lower() for p in prefixes)): return " ".join(b)
        return ""
    t11 = find("Tez Çalışması Durumu", "Thesis Status", "Thesis Work Status")
    t22 = find("Seminer - Tez", "Seminer – Tez", "Seminer Tez", "Seminar - Thesis", "Seminar – Thesis", "Seminar–Thesis", "Seminar Thesis")
    # alt sınırlar (form.md)
    wc = lambda t: len(t.split())
    mins = [("Tez Çalışması Durumu", 150, t11), ("Tezin … ile İlişkisi", 120, find("Tezin", "Relation of the Thesis")),
            ("Seminer Konusu", 200, find("Seminer Konusu", "Seminar Topic")), ("Seminer – Tez İlişkisi", 60, t22)]
    for name, lim, txt in mins:
        if txt and wc(txt) < lim: out.append(f"'{name}' {wc(txt)} sözcük; en az {lim}.")
    abstract_txt = next((t for st, t in paras if st == "Abstract"), "")
    body_total = sum(wc(" ".join(b)) for h, b in sections.items() if not is_refs(h) and not is_tesekkur(h)) + wc(abstract_txt)
    if body_total < 850: out.append(f"Form metni {body_total} sözcük (özet dahil, kaynaklar ve teşekkür hariç); en az 850, hedef 1.000+ (2–3 sayfa). Tek sayfalık form yetersizdir.")
    if t11 and t22:
        r = difflib.SequenceMatcher(None, t11, t22).ratio()
        if r > 0.6: out.append(f"1.1 ile 2.2 metinleri %{r*100:.0f} benzer; 2.2 seminerin tezdeki yerini anlatmalı.")
    if not any(is_beyan(h) for h in sections): out.append("Yapay zekâ kullanım beyanı bölümü yok.")
    abstract = next((t for st, t in paras if st == "Abstract"), "")
    n = len(abstract.split())
    if abstract and not (120 <= n <= 280): out.append(f"Özet {n} sözcük; hedef 150–250.")
    for st, t in paras:
        if st != "references" and t:
            # yer tutucunun içindeki LaTeX zaten "işaret kaldı" ile bildirilir
            for h in math_hits(MARK.sub("", t)): out.append(f"Elle yazılmış matematik, Word denklemine çevirin ({h})")
    refs = [t for st, t in paras if st == "references" and t]
    if len(refs) < 5: out.append(f"{len(refs)} referans; en az 5 beklenir.")
    hakemli = sum(1 for r in refs if re.search(r"\bc\.\s*\d|vol\.|doi:|Proc\.|Conf\.|Konferans|Sempozyum|Forum|J\.|Journal|Dergi|Transactions|Letters", r))
    if hakemli < 3: out.append(f"Hakemli makale ya da bildiri görünen referans {hakemli}; en az 3 (dergi/konferans adı, cilt ya da DOI ile).")
    for i, r in enumerate(refs, 1):
        if not re.search(r"(19|20)\d\d", r): out.append(f"Referans [{i}] yıl içermiyor.")
    cited = set()
    for m in re.finditer(r"\[(\d+)(?:\]\s*[–-]\s*\[(\d+))?\]", full):
        a = int(m.group(1)); b = int(m.group(2)) if m.group(2) else a
        cited.update(range(a, b + 1))
    cited.discard(0)
    for k in sorted(cited):
        if k > len(refs): out.append(f"Metinde [{k}] anılıyor ama {len(refs)} referans var.")
    for k in range(1, len(refs) + 1):
        if k not in cited: out.append(f"Referans [{k}] metinde hiç anılmamış.")
    return out

def shrink_ratio(pic):
    """Görselin çizildiği boyuta göre küçültme oranı; çözünürlük bilgisi yoksa None."""
    dx, _ = pic.image.dpi
    if dx == 72:          # python-pptx'in varsayılanı: bilgi yok
        return None
    return pic.width / 914400 / (pic.image.size[0] / dx)

def check_visuals(idx, title, s, out, chart_sizes):
    """Görsel sayısı, küçültme, küçük görsel; PowerPoint grafiklerinde seri sayısı ve punto."""
    from pptx.util import Pt
    visuals = []
    def walk(sh):
        if sh.shape_type == 6:
            for c in sh.shapes: walk(c)
        elif sh.shape_type == 13 or getattr(sh, "has_chart", False):
            visuals.append(sh)
    for sh in s.shapes: walk(sh)
    # tıklamayla üst üste açılan katmanlar aynı merkezdedir; aynı anda görünen görsel sayısı = farklı merkez sayısı
    centers = {(round((v.left + v.width / 2) / 914400, 1), round((v.top + v.height / 2) / 914400, 1)) for v in visuals}
    if len(centers) > 3:
        out.append(f"Slayt {idx} ({title[:30]}): aynı anda {len(centers)} görsel/grafik görünüyor; 3'ten fazlası okunmaz. Bölün ya da mesajı taşıyanları seçin (okunabilirlik önerisi).")
    small, shrunk = [], []
    for v in visuals:
        w_in, h_in = v.width / 914400, v.height / 914400
        if w_in < 2.5 and h_in < 2.5:
            small.append(f"{w_in:.1f} × {h_in:.1f}")
        if v.shape_type == 13:
            oran = shrink_ratio(v)
            if oran is not None and oran < 0.7:
                shrunk.append((oran, f"{w_in:.1f} × {h_in:.1f}"))
    if small:
        out.append(f"Slayt {idx} ({title[:30]}): {len(small)} görsel çok küçük ({', '.join(sorted(set(small)))} inç); grafikse içindeki yazı okunmaz.")
    if shrunk:
        en = min(shrunk)
        out.append(f"Slayt {idx} ({title[:30]}): {len(shrunk)} görsel çizildiği boyuttan küçültülmüş (en çok %{en[0] * 100:.0f}'ine); içindeki yazı da küçüldü (16 pt → ≈ {16 * en[0]:.0f} pt). Slayttaki boyutunda ({en[1]} inç) yeniden üretin.")
        if getattr(v, "has_chart", False):
            ch = v.chart
            n = sum(len(list(pl.series)) for pl in ch.plots)
            if n > 4:
                out.append(f"Slayt {idx} ({title[:30]}): grafikte {n} seri; 4'ten fazlası okunmaz. Asıl seriyi öne çıkarın ya da grafiği bölün (okunabilirlik önerisi).")
            sizes = set()
            if ch.font.size is not None: sizes.add(ch.font.size.pt)
            for ax_name in ("category_axis", "value_axis"):
                try:
                    ax = getattr(ch, ax_name)
                    if ax.tick_labels.font.size is not None: sizes.add(ax.tick_labels.font.size.pt)
                except Exception:
                    pass
            if ch.has_legend and ch.legend.font.size is not None: sizes.add(ch.legend.font.size.pt)
            small = sorted(z for z in sizes if z < 14)
            if small:
                out.append(f"Slayt {idx} ({title[:30]}): grafik yazısı {', '.join(f'{z:g}' for z in small)} pt; slaytta en az 14 pt olmalı.")
            for z in sizes: chart_sizes.setdefault(z, []).append(idx)

def check_pptx(path):
    from pptx import Presentation
    from pptx.util import Pt
    out = []
    prs = Presentation(path)
    slides = list(prs.slides)
    titles = []
    all_text = []
    title_sizes = {}           # punto -> [slayt]
    chart_sizes = {}           # PowerPoint grafik nesnelerinin puntoları -> [slayt]
    body_sizes = {}            # punto -> [slayt]
    for idx, s in enumerate(slides, 1):
        title = ""
        for sh in s.shapes:
            if sh.has_text_frame and sh.name == "baslik":
                title = sh.text_frame.text.strip()
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.size is not None and r.text.strip(): title_sizes.setdefault(r.font.size.pt, []).append(idx)
        if not title and idx == 1: title = "(kapak)"
        if not title:
            for sh in s.shapes:
                if sh.name in ("kapanis_baslik",): title = "Teşekkürler"
        titles.append(title or f"(slayt {idx})")
        texts = []
        def walk(shape):
            if shape.shape_type == 6:
                for c in shape.shapes: walk(c)
            elif shape.has_text_frame:
                texts.append(shape)
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        t = cell.text_frame.text
                        all_text.append(t)
                        for m in MARK.finditer(t): out.append(f"Slayt {idx}: işaret kaldı (tablo): {m.group(0)[:50]}")
                        for p in cell.text_frame.paragraphs:
                            for r in p.runs:
                                if r.font.size is not None and r.font.size < Pt(14) and r.text.strip():
                                    out.append(f"Slayt {idx} ({title[:30]}): tabloda {r.font.size.pt:.0f} pt metin"); break
        for sh in s.shapes: walk(sh)
        if idx != 1 and not has_k(title, K_TESEKKUR):
            check_visuals(idx, title, s, out, chart_sizes)
        bullets = 0
        slide_sizes = set()
        for sh in texts:
            t = sh.text_frame.text
            all_text.append(t)
            for m in MARK.finditer(t): out.append(f"Slayt {idx}: işaret kaldı: {m.group(0)[:50]}")
            if not has_k(title, K_KAYNAK) and sh.name not in ("altbilgi", "sayfa_no"):
                for line in t.split("\n"):
                    for h in math_hits(MARK.sub("", line)): out.append(f"Slayt {idx} ({title[:30]}): elle yazılmış matematik, denklem aracıyla yazın ({h})")
            if sh.name in ("baslik", "altbilgi", "sayfa_no", "tesekkur_satirlar", "kapak_kunye", "tablo_not") or sh.name.startswith(("altyazi", "sayi", "gorsel", "zaman", "mesaj", "kapak_", "kapanis_")) or sh.name.endswith(("_no", "_baslik")) or has_k(title, K_KAYNAK): continue
            sh_bullets = 0
            for p in sh.text_frame.paragraphs:
                if not p.text.strip(): continue
                pPr = p._p.pPr
                if pPr is not None and pPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar") is not None: sh_bullets += 1
                for r in p.runs:
                    if r.font.size is not None and r.text.strip() and not sh.name.startswith("adim") and not has_k(title, K_ICINDEKILER):
                        slide_sizes.add(r.font.size.pt)
                for r in p.runs:
                    if r.font.size is not None and r.font.size < Pt(18) and r.text.strip():
                        out.append(f"Slayt {idx} ({title[:30]}): {r.font.size.pt:.0f} pt metin: \"{r.text[:30]}\""); break
            bullets = max(bullets, sh_bullets)
        if len(slide_sizes) > 1:
            out.append(f"Slayt {idx} ({title[:30]}): aynı slaytta farklı gövde puntoları ({', '.join(f'{z:g}' for z in sorted(slide_sizes))}); vurgu için kalın yazı kullanın, puntoyu değiştirmeyin.")
        for z in slide_sizes: body_sizes.setdefault(z, []).append(idx)
        if bullets > 5: out.append(f"Slayt {idx} ({title[:30]}): bir kutuda {bullets} madde; en çok 5.")
        if not s.has_notes_slide or not s.notes_slide.notes_text_frame.text.strip():
            if idx != 1 and not has_k(title, K_YARDIMCI):
                out.append(f"Slayt {idx} ({title[:30]}): konuşma notu yok.")
    # deste geneli punto tutarlılığı: en sık punto esas, farklı olan slaytlar listelenir
    for ad, sizes in (("Başlık", title_sizes), ("Gövde", body_sizes), ("Grafik yazısı", chart_sizes)):
        if len(sizes) > 1:
            esas = max(sizes, key=lambda z: len(sizes[z]))
            farkli = "; ".join(f"{z:g} pt: slayt {', '.join(map(str, sorted(set(v))))}" for z, v in sorted(sizes.items()) if z != esas)
            out.append(f"{ad} puntosu destede tutarsız: çoğu slayt {esas:g} pt, farklı olanlar {farkli}.")
    # görselsiz içerik slaytı oranı
    YAPI = ("gorsel", "tablo", "kart", "karsi", "adim", "zaman")
    def has_pic(shape):
        if shape.shape_type == 13: return True
        if shape.shape_type == 6: return any(has_pic(c) for c in shape.shapes)
        if shape.name.startswith(YAPI) and "[ÖĞRENCİ TAMAMLAYACAK" not in (shape.text_frame.text if shape.has_text_frame else ""): return True
        return False
    text_only = 0; content_n = 0
    for idx, (s, t) in enumerate(zip(slides, titles), 1):
        if idx == 1 or has_k(t, K_YARDIMCI): continue
        content_n += 1
        has_eq = s._element.find(".//{http://schemas.microsoft.com/office/drawing/2010/main}m") is not None   # PowerPoint denklemi
        if not has_eq and not any(has_pic(sh) for sh in s.shapes): text_only += 1
    if content_n and text_only / content_n > 0.34:
        out.append(f"{content_n} içerik slaydının {text_only} tanesi salt madde listesi (üst sınır üçte bir); tablo, süreç, karşılaştırma, kart ya da görsel kalıbına çevirin.")
    icerik = [t for t in titles[1:] if not has_k(t, K_YARDIMCI)]
    if not (15 <= len(icerik) <= 20): out.append(f"İçerik slaydı sayısı {len(icerik)}; 15–20 önerilir (toplam {len(slides)}).")
    if not any(has_k(t, K_ICINDEKILER) for t in titles): out.append("İçindekiler slaydı yok.")
    if not any(has_k(t, K_KAYNAK) for t in titles): out.append("Kaynaklar slaydı yok (zorunlu).")
    if not any(has_k(t, K_TESEKKUR) for t in titles): out.append("Teşekkürler slaydı yok (zorunlu).")
    joined = "\n".join(all_text)
    if not re.search(r"yapay zek|\bAI\b|artificial intelligence|language model|claude|chatgpt|gpt|gemini|copilot", joined, re.I): out.append("Yapay zekâ beyanı satırı sunumda yok (Teşekkürler slaydı).")
    # kaynak numaraları
    ref_count = 0
    for s, t in zip(slides, titles):
        if has_k(t, K_KAYNAK):
            for sh in s.shapes:
                if sh.has_text_frame and sh.name == "govde":
                    ref_count += sum(1 for p in sh.text_frame.paragraphs if p.text.strip())
    cited = set()
    for m in re.finditer(r"\[(\d+)(?:\]\s*[–-]\s*\[(\d+))?\]", joined):
        a = int(m.group(1)); b = int(m.group(2)) if m.group(2) else a
        cited.update(range(a, b + 1))
    for k in sorted(cited):
        if ref_count and k > ref_count: out.append(f"Slaytlarda [{k}] anılıyor ama Kaynaklar'da {ref_count} kayıt var.")
    return out

def main():
    if len(sys.argv) < 2:
        print(__doc__); sys.exit(2)
    problems = 0
    for path in sys.argv[1:]:
        if path.lower().endswith(".docx"): res = check_docx(path)
        elif path.lower().endswith(".pptx"): res = check_pptx(path)
        else: print(f"{path}: desteklenmeyen uzantı"); continue
        res = list(dict.fromkeys(res))   # tekrarları at, sırayı koru
        print(f"== {path}: {'sorun yok' if not res else str(len(res)) + ' uyarı'}")
        for r in res: print("  -", r)
        problems += len(res)
    sys.exit(1 if problems else 0)

if __name__ == "__main__":
    main()
